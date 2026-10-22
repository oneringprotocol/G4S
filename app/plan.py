"""G4S route plan (Sheet1) and Geofence Directory.

Sheet1 has a vehicle table at the top, then one block per vehicle:
    DRIVER NAME | Company | MOB NO | Vehicle | Emts | Type | nth Trip (From, To, staff) x N
    outbound row   (From = accommodation, To = sites, newline separated)
    return row     (From = sites, To = accommodation)
    'MSS <branch>' row
    times row      (start, end) per trip
    'Hours worked' row
"""
import difflib
import json
import re
from dataclasses import dataclass, field, asdict
from datetime import time
from typing import Dict, List, Optional

import openpyxl

from . import config
from .geo import clean, dist_m, parse_coords


@dataclass
class Site:
    id: str
    name: str          # as in the directory
    label: str         # client wording
    lat: float
    lon: float
    accommodation: bool
    radius: int


@dataclass
class PlannedTrip:
    n: int
    group: str
    start: str         # "HH:MM" local
    end: str
    sites: List[str]   # site ids, in plan order
    staff_out: Optional[int] = None
    staff_back: Optional[int] = None
    unresolved: List[str] = field(default_factory=list)


@dataclass
class Vehicle:
    vehicle: str
    driver: str
    type: str
    type_abbr: str
    branch: str
    emirates: List[str]
    home: Optional[str]           # site id of the home accommodation
    trips: List[PlannedTrip]
    phone: Optional[str] = None


def _norm(s):
    s = str(s or "").lower()
    s = re.sub(r"accomm?odi?ation", "accommodation", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def load_sites(wb) -> List[Site]:
    over = json.loads((config.DATA / "site_overrides.json").read_text())
    seen, sites = set(), []

    def add(name, lat, lon):
        key = (_norm(name), round(lat, 4), round(lon, 4))
        if key in seen:
            return
        seen.add(key)
        acc = "accommodation" in _norm(name)
        sid = "s%d" % len(sites)
        sites.append(Site(sid, " ".join(str(name).split()), clean(name), lat, lon, acc,
                          config.ACCOMMODATION_M if acc else config.GEOFENCE_M))

    ws = wb[config.GEOFENCE_TAB]
    for row in ws.iter_rows(values_only=True):
        name, coords = row[0], row[1] if len(row) > 1 else None
        if not name or name == "Site Name":
            continue
        fix = over["set"].get(" ".join(str(name).split()))
        ll = tuple(fix) if fix else parse_coords(coords)
        if ll:
            add(name, *ll)
    for name, ll in over["add"].items():
        add(name, *ll)
    return sites


class SiteIndex:
    def __init__(self, sites: List[Site]):
        self.sites = sites
        self.by_id = {s.id: s for s in sites}
        self.by_norm: Dict[str, List[Site]] = {}
        for s in sites:
            self.by_norm.setdefault(_norm(s.name), []).append(s)

    def find(self, name, near=None) -> Optional[Site]:
        """Directory site for a name; when a name exists twice, the one nearest `near`."""
        n = _norm(name)
        cands = self.by_norm.get(n)
        if not cands:
            close = difflib.get_close_matches(n, self.by_norm.keys(), n=1, cutoff=0.86)
            cands = self.by_norm[close[0]] if close else None
        if not cands:
            return None
        if near and len(cands) > 1:
            return min(cands, key=lambda s: dist_m(near[0], near[1], s.lat, s.lon))
        return cands[0]

    def nearest(self, lat, lon):
        best = min(self.sites, key=lambda s: dist_m(lat, lon, s.lat, s.lon))
        return best, dist_m(lat, lon, best.lat, best.lon)

    def inside(self, lat, lon):
        """Nearest site whose geofence contains the point, else None."""
        hit, hd = None, None
        for s in self.sites:
            d = dist_m(lat, lon, s.lat, s.lon)
            if d <= s.radius and (hd is None or d < hd):
                hit, hd = s, d
        return hit


def _t(v):
    if isinstance(v, time):
        return v.strftime("%H:%M")
    m = re.match(r"\s*(\d{1,2})\s*:\s*(\d{2})\s*(AM|PM)?", str(v or ""), re.I)
    if not m:
        return None
    h = int(m.group(1)) % 12 + (12 if (m.group(3) or "").upper() == "PM" else 0) if m.group(3) else int(m.group(1))
    return "%02d:%s" % (h, m.group(2))


def _names(cell):
    if cell is None:
        return []
    return [p.strip() for p in str(cell).split("\n") if p.strip()]


def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


_EMIRATES = {"dubai": "Dubai", "sharjah": "Sharjah", "shj": "Sharjah", "rak": "Ras Al Khaimah",
             "fuj": "Fujairah", "auh": "Abu Dhabi", "abu dhabi": "Abu Dhabi", "ajman": "Ajman"}


def _emirates(*branches):
    out = []
    for b in branches:
        for part in re.split(r"[/,]", str(b or "")):
            e = _EMIRATES.get(part.strip().lower())
            if e and e not in out:
                out.append(e)
    return out


def load_plan():
    wb = openpyxl.load_workbook(config.PLAN_XLSX, data_only=True)
    sites = SiteIndex(load_sites(wb))
    aliases = json.loads((config.DATA / "plan_aliases.json").read_text())
    rows = [list(r) for r in wb[config.PLAN_TAB].iter_rows(values_only=True)]
    wb.close()

    # vehicle table
    info = {}
    for i, r in enumerate(rows):
        if len(r) > 1 and r[1] == "New Numbers":
            for r2 in rows[i + 1:]:
                if not isinstance(r2[1], (int, float)):
                    break
                info[str(int(r2[1]))] = {"type": r2[3], "driver": r2[6], "branch": r2[9]}
            break

    vehicles: Dict[str, Vehicle] = {}
    for i, r in enumerate(rows):
        if len(r) < 2 or r[1] != "DRIVER NAME":
            continue
        head = rows[i]
        trip_cols = [c for c, v in enumerate(head) if isinstance(v, str) and "Trip" in v]
        data, times = [], None
        for r2 in rows[i + 2:]:
            if len(r2) > 4 and isinstance(r2[4], (int, float)):
                data.append(r2)
            elif any(_t(v) for v in r2[7:9] if isinstance(v, time)):
                times = r2
                break
        if not data:
            continue
        out, back = data[0], data[1] if len(data) > 1 else [None] * len(data[0])
        vno = str(int(out[4]))
        meta = info.get(vno, {})
        vt = str(out[6] or meta.get("type") or "")
        al = dict(aliases.get("*", {}), **aliases.get(vno, {}))

        def resolve(name, near=None):
            key = next((k for k in al if _norm(k) == _norm(name)), None)
            names = al[key] if key else [name]
            hits, miss = [], []
            for nm in names:
                s = sites.find(nm, near)
                (hits.append(s.id) if s else miss.append(nm))
            return hits, miss

        home_name = _names(out[trip_cols[0]])[0] if trip_cols else None
        home_ids, _ = resolve(home_name) if home_name else ([], [])
        home = home_ids[0] if home_ids else None
        near = (sites.by_id[home].lat, sites.by_id[home].lon) if home else None

        trips = []
        for k, c in enumerate(trip_cols):
            to_names = _names(out[c + 1] if c + 1 < len(out) else None)
            back_names = _names(back[c] if c < len(back) else None)
            start = _t(times[c]) if times and c < len(times) else None
            end = _t(times[c + 1]) if times and c + 1 < len(times) else None
            if not (to_names or back_names) or not start:
                continue
            ids, miss = [], []
            for nm in to_names + back_names:
                h, m = resolve(nm, near)
                ids += [x for x in h if x not in ids and x != home]
                miss += [x for x in m if x not in miss]
            trips.append(PlannedTrip(
                n=len(trips) + 1,
                group=clean(" / ".join(to_names) if to_names and _norm(to_names[0]) != _norm(home_name) else " / ".join(back_names)),
                start=start, end=end, sites=ids,
                staff_out=_int(out[c + 2] if c + 2 < len(out) else None),
                staff_back=_int(back[c + 2] if c + 2 < len(back) else None),
                unresolved=miss))
        vehicles[vno] = Vehicle(
            vehicle=vno,
            driver=clean(out[1] or meta.get("driver")),
            type=vt,
            type_abbr="".join(w[0] for w in vt.split()).upper()[:2] if " " in vt else {"dongfeng": "DF"}.get(vt.lower(), vt[:2].upper()),
            branch=clean(out[5] or meta.get("branch")),
            emirates=_emirates(out[5], meta.get("branch")),
            home=home, trips=trips,
            phone=str(out[3]) if out[3] else None)
    return vehicles, sites


def as_dict(obj):
    return asdict(obj)
