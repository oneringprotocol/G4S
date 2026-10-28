"""Operations log: the permanent vehicle directory and the replacement log.

These replace the first two tabs of the Internal G4S Operations Sheet
("Permanent Vehicles - Internal" and "Replacement Vehicles"). The ops team
logs new entries in the Operations tab; every change is kept in ops_history.

Replacement entries also feed the route analysis: while a permanent vehicle
is replaced, its planned trips show as "Run by replacement" (records()).

One time import of the sheet history:
    python3 -m app.ops import "sheets/Internal_ G4S Operations Sheet.xlsx"
"""
import json
import re
import sqlite3
import sys
import time
from datetime import date, datetime, timedelta

from . import config

STATUSES = [
    "Active", "Active - Out Source", "Returned", "Returned - Out Source",
    "Replacement Not Provided", "Reported To The Workshop - No Replacement",
    "Vehicle Not Recovered - Replacement Not Provided", "Returned From Workshop - Fixed",
    "Permanent Vehicle In Workshop - Replacement Vehicle Also Break Down",
]
ACTIVE = ("Active", "Active - Out Source")      # a replacement vehicle is running
NO_REPLACEMENT = ("Replacement Not Provided", "Reported To The Workshop - No Replacement",
                  "Vehicle Not Recovered - Replacement Not Provided",
                  "Permanent Vehicle In Workshop - Replacement Vehicle Also Break Down")
OPEN = ACTIVE + NO_REPLACEMENT                    # the permanent vehicle is not back yet
EMIRATES = ["Abu Dhabi", "Dubai", "Sharjah", "Ajman", "RAK", "Fujairah", "Umm Al Quwain", "Al Ain"]
VEHICLE_TYPES = ["5 Seater", "12 Seater", "33 Seater", "Hilux"]
VEHICLE_STATUSES = ["Brand New", "Pre Lease", "Permanent replacement of total Loss"]
MAIN_SUPPLIER = "AGT"

VEHICLE_FIELDS = ["plate", "swvl_no", "supplier", "agt_agreement", "g4s_agreement", "plate_code", "vehicle_type",
                  "emirate", "driver", "contact", "delivery_date", "delivery_time", "start_date", "vehicle_status",
                  "rental_charges", "notes"]
# swvl_no ... contact are copied from the directory when the entry is logged, so old rows keep
# the driver and emirate of that time even after the directory changes (as the sheet values do)
REPLACEMENT_FIELDS = ["plate", "swvl_no", "vehicle_status", "vehicle_type", "emirate", "driver", "contact",
                      "replacement", "supplier", "replacement_date", "status", "return_date", "workshop_date",
                      "perm_return_date", "issue", "issue_date", "remarks"]
SNAPSHOT = ["swvl_no", "vehicle_status", "vehicle_type", "emirate", "driver", "contact"]
DATES = {"delivery_date", "start_date", "replacement_date", "return_date", "workshop_date", "perm_return_date", "issue_date"}
TABLES = {"vehicles": ("ops_vehicles", VEHICLE_FIELDS), "replacements": ("ops_replacements", REPLACEMENT_FIELDS)}


class OpsError(Exception):
    def __init__(self, msg, code=400):
        super().__init__(msg)
        self.code = code


def _db():
    con = sqlite3.connect(config.DB_FILE, check_same_thread=False)
    con.row_factory = sqlite3.Row
    return con


def init():
    cols = lambda fs: ", ".join("%s TEXT NOT NULL DEFAULT ''" % f for f in fs if f != "plate")
    with _db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS ops_vehicles(
            id INTEGER PRIMARY KEY, plate TEXT UNIQUE NOT NULL, %s,
            created_by TEXT, created_at INTEGER, updated_by TEXT, updated_at INTEGER);
        CREATE TABLE IF NOT EXISTS ops_replacements(
            id INTEGER PRIMARY KEY, plate TEXT NOT NULL, %s,
            created_by TEXT, created_at INTEGER, updated_by TEXT, updated_at INTEGER);
        CREATE TABLE IF NOT EXISTS ops_history(
            id INTEGER PRIMARY KEY, at INTEGER NOT NULL, user TEXT NOT NULL, tab TEXT NOT NULL,
            row_id INTEGER NOT NULL, action TEXT NOT NULL, changes TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS ops_history_row ON ops_history(tab, row_id);
        """ % (cols(VEHICLE_FIELDS), cols(REPLACEMENT_FIELDS)))


# ---------- values ----------
def plate_of(s):
    """'29357 (New Vehicle 55045)' -> '29357'."""
    m = re.match(r"\s*(\d+)", str(s or ""))
    return m.group(1) if m else ""


def _iso(v):
    return isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) is not None


def _clean(tab, body):
    fields = TABLES[tab][1]
    out = {}
    for f in fields:
        v = body.get(f, "")
        v = "" if v is None else " ".join(str(v).split())
        if f in DATES and v and not _iso(v):
            raise OpsError("%s must be a date (YYYY-MM-DD)." % f.replace("_", " ").capitalize())
        out[f] = v
    if not out["plate"]:
        raise OpsError("The plate number of the permanent vehicle is required." if tab == "replacements"
                       else "The plate number is required.")
    if tab == "replacements":
        if out["status"] not in STATUSES:
            raise OpsError("Choose a status from the list.")
        if out["return_date"] and out["replacement_date"] and out["return_date"] < out["replacement_date"]:
            raise OpsError("The return date is before the replacement date.")
        if out["perm_return_date"] and out["workshop_date"] and out["perm_return_date"] < out["workshop_date"]:
            raise OpsError("The permanent vehicle return date is before the workshop report date.")
    return out


def replacement_type(r):
    """Column P of the sheet."""
    if r["status"] in ("Vehicle Not Recovered - Replacement Not Provided", "Reported To The Workshop - No Replacement"):
        return "Reported To Workshop - No Replacement"
    if not r["replacement"] and not r["supplier"]:
        return "Replacement Was Not Arranged"
    return "Provided By AGT" if r["supplier"].upper() == MAIN_SUPPLIER else "Out Source Replacement"


def workshop_days(r, today=None):
    """Column O of the sheet: workshop report date to permanent vehicle return date (or today)."""
    if not _iso(r["workshop_date"]):
        return None
    end = r["perm_return_date"] if _iso(r["perm_return_date"]) else (today or date.today().isoformat())
    return (date.fromisoformat(end) - date.fromisoformat(r["workshop_date"])).days


# ---------- read ----------
def _rows(con, table):
    return [dict(r) for r in con.execute("SELECT * FROM %s ORDER BY id" % table)]


def snapshot():
    """Everything the Operations tab shows."""
    today = date.today().isoformat()
    with _db() as con:
        vehicles, reps = _rows(con, "ops_vehicles"), _rows(con, "ops_replacements")
    for r in reps:
        r["workshop_days"] = workshop_days(r, today)
        r["replacement_type"] = replacement_type(r)
        r["open"] = r["status"] in OPEN
    open_by = {}
    for r in reps:
        if r["open"]:
            open_by.setdefault(plate_of(r["plate"]), []).append(r["id"])
    count_by = {}
    for r in reps:
        count_by[plate_of(r["plate"])] = count_by.get(plate_of(r["plate"]), 0) + 1
    for v in vehicles:
        v["open"] = open_by.get(v["plate"], [])
        v["replacements"] = count_by.get(v["plate"], 0)
    suppliers = sorted({r["supplier"] for r in reps if r["supplier"]} | {MAIN_SUPPLIER}, key=lambda s: (s != MAIN_SUPPLIER, s.lower()))
    issues = sorted({r["issue"] for r in reps if r["issue"]}, key=str.lower)
    return {"vehicles": vehicles, "replacements": reps, "summary": summary(reps, vehicles),
            "options": {"statuses": STATUSES, "active": list(ACTIVE), "no_replacement": list(NO_REPLACEMENT),
                        "emirates": sorted(set(EMIRATES) | {v["emirate"] for v in vehicles if v["emirate"]}),
                        "vehicle_types": sorted(set(VEHICLE_TYPES) | {v["vehicle_type"] for v in vehicles if v["vehicle_type"]}),
                        "vehicle_statuses": sorted(set(VEHICLE_STATUSES) | {v["vehicle_status"] for v in vehicles if v["vehicle_status"]}),
                        "suppliers": suppliers, "main_supplier": MAIN_SUPPLIER, "issues": issues},
            "today": today}


def summary(reps, vehicles):
    act = [r for r in reps if r["status"] == "Active"]
    out = [r for r in reps if r["status"] == "Active - Out Source"]
    waiting = [r for r in reps if r["status"] in NO_REPLACEMENT]
    days = [r["workshop_days"] for r in reps if r["open"] and r["workshop_days"] is not None]
    by_em = {}
    for r in act + out:
        e = by_em.setdefault(r["emirate"] or "Unknown", {"emirate": r["emirate"] or "Unknown", "replacement": 0, "outsource": 0})
        e["replacement" if r["status"] == "Active" else "outsource"] += 1
    month = date.today().isoformat()[:7]
    return {"active": len(act) + len(out), "replacement": len(act), "outsource": len(out),
            "waiting": len(waiting), "in_workshop": sum(1 for r in reps if r["open"]),
            "avg_days": round(sum(days) / len(days)) if days else None, "max_days": max(days) if days else None,
            "permanent": len(vehicles),
            "by_type": {t: sum(1 for v in vehicles if v["vehicle_type"] == t) for t in sorted({v["vehicle_type"] for v in vehicles})},
            "by_emirate": sorted(by_em.values(), key=lambda e: -(e["replacement"] + e["outsource"])),
            "this_month": sum(1 for r in reps if r["replacement_date"][:7] == month or r["issue_date"][:7] == month)}


def history(tab, row_id):
    if tab not in TABLES:
        raise OpsError("Unknown table.", 404)
    with _db() as con:
        return [dict(r, changes=json.loads(r["changes"])) for r in con.execute(
            "SELECT at, user, action, changes FROM ops_history WHERE tab=? AND row_id=? ORDER BY id DESC", (tab, row_id))]


def records():
    """Maintenance records for the route analysis, by plate: the days a replacement ran instead of
    the permanent vehicle (replacement date up to the day before the return date, open while active)."""
    out = {}
    with _db() as con:
        rows = con.execute("SELECT plate, replacement, replacement_date, return_date, status, issue FROM ops_replacements "
                           "WHERE replacement != '' AND replacement_date != '' ORDER BY replacement_date").fetchall()
    for r in rows:
        plate, by = plate_of(r["plate"]), plate_of(r["replacement"]) or r["replacement"]
        if not plate or not _iso(r["replacement_date"]):
            continue
        if _iso(r["return_date"]):
            to = max(r["replacement_date"], (date.fromisoformat(r["return_date"]) - timedelta(days=1)).isoformat())
        elif r["status"] in ACTIVE:
            to = "9999-12-31"
        else:
            to = r["replacement_date"]   # returned, but the return date was never logged
        out.setdefault(plate, []).append({"from": r["replacement_date"], "to": to, "note": r["issue"],
                                          "replaced_by": by, "source": "ops"})
    return out


# ---------- write ----------
def _fill_snapshot(con, row):
    """New replacement entries copy the vehicle details from the directory where left empty."""
    v = con.execute("SELECT * FROM ops_vehicles WHERE plate=?", (plate_of(row["plate"]),)).fetchone()
    if v:
        for f in SNAPSHOT:
            if not row[f]:
                row[f] = v[f]


def _log(con, user, tab, row_id, action, changes):
    con.execute("INSERT INTO ops_history(at, user, tab, row_id, action, changes) VALUES(?,?,?,?,?,?)",
                (int(time.time()), user, tab, row_id, action, json.dumps(changes)))


def create(tab, body, user):
    table, fields = TABLES[tab]
    row = _clean(tab, body)
    now = int(time.time())
    with _db() as con:
        if tab == "replacements":
            _fill_snapshot(con, row)
        elif con.execute("SELECT 1 FROM ops_vehicles WHERE plate=?", (row["plate"],)).fetchone():
            raise OpsError("Plate %s is already in the directory." % row["plate"], 409)
        # never reuse the id of a deleted entry, or the new one would show its change history
        rid = max(con.execute("SELECT COALESCE(MAX(id), 0) FROM %s" % table).fetchone()[0],
                  con.execute("SELECT COALESCE(MAX(row_id), 0) FROM ops_history WHERE tab=?", (tab,)).fetchone()[0]) + 1
        con.execute("INSERT INTO %s(id, %s, created_by, created_at, updated_by, updated_at) VALUES(%s)"
                    % (table, ", ".join(fields), ", ".join("?" * (len(fields) + 5))),
                    [rid] + [row[f] for f in fields] + [user, now, user, now])
        _log(con, user, tab, rid, "created", {f: [None, row[f]] for f in fields if row[f]})
        return dict(con.execute("SELECT * FROM %s WHERE id=?" % table, (rid,)).fetchone())


def update(tab, row_id, body, user):
    table, fields = TABLES[tab]
    row = _clean(tab, body)
    with _db() as con:
        old = con.execute("SELECT * FROM %s WHERE id=?" % table, (row_id,)).fetchone()
        if not old:
            raise OpsError("This entry no longer exists. Reload the page.", 404)
        if body.get("updated_at") is not None and int(body["updated_at"]) != (old["updated_at"] or 0):
            raise OpsError("%s changed this entry while you were editing. Close it and open it again to see the latest."
                           % (old["updated_by"] or "Someone"), 409)
        if tab == "vehicles" and row["plate"] != old["plate"] and \
                con.execute("SELECT 1 FROM ops_vehicles WHERE plate=? AND id!=?", (row["plate"], row_id)).fetchone():
            raise OpsError("Plate %s is already in the directory." % row["plate"], 409)
        changes = {f: [old[f], row[f]] for f in fields if (old[f] or "") != row[f]}
        if changes:
            con.execute("UPDATE %s SET %s, updated_by=?, updated_at=? WHERE id=?"
                        % (table, ", ".join("%s=?" % f for f in fields)),
                        # always moves on, so two saves in the same second still count as a change
                        [row[f] for f in fields] + [user, max(int(time.time()), (old["updated_at"] or 0) + 1), row_id])
            _log(con, user, tab, row_id, "edited", changes)
        return dict(con.execute("SELECT * FROM %s WHERE id=?" % table, (row_id,)).fetchone())


def delete(tab, row_id, user):
    table, fields = TABLES[tab]
    with _db() as con:
        old = con.execute("SELECT * FROM %s WHERE id=?" % table, (row_id,)).fetchone()
        if not old:
            raise OpsError("This entry no longer exists.", 404)
        con.execute("DELETE FROM %s WHERE id=?" % table, (row_id,))
        _log(con, user, tab, row_id, "deleted", {f: [old[f], None] for f in fields if old[f]})


# ---------- one time import from the sheet ----------
def _cell(v, field=""):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%H:%M") if field.endswith("time") else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if hasattr(v, "hour") and hasattr(v, "minute"):   # datetime.time
        return "%02d:%02d" % (v.hour, v.minute)
    if isinstance(v, float):
        v = int(v) if v.is_integer() else round(v, 2)
    s = " ".join(str(v).split())
    return "" if s in ("-", "#N/A", "#REF!", "#VALUE!") else s


def _phone(s):
    return "0" + s if re.fullmatch(r"5\d{8}", s) else s   # 554002387 -> 0554002387


def import_xlsx(path, replace=False):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    init()
    with _db() as con:
        if not replace and (con.execute("SELECT 1 FROM ops_vehicles").fetchone() or
                            con.execute("SELECT 1 FROM ops_replacements").fetchone()):
            raise OpsError("The operations tables already have data. Use --replace to overwrite them.")
        con.execute("DELETE FROM ops_vehicles")
        con.execute("DELETE FROM ops_replacements")
        now, who = int(time.time()), "Imported history"

        rows = list(wb["Permanent Vehicles - Internal"].iter_rows(values_only=True))
        head = rows[0]
        main = ["supplier", "swvl_no", "agt_agreement", "g4s_agreement", "plate_code", "plate", "vehicle_type", "emirate",
                "driver", "contact", "delivery_date", "delivery_time", "start_date", "vehicle_status", "rental_charges"]
        nv = 0
        for r in rows[1:]:
            if not r or not _cell(r[5]):
                continue
            v = {f: _cell(r[i], f) for i, f in enumerate(main)}
            v["contact"] = _phone(v["contact"])
            # columns P to Z hold pre lease and permanent replacement notes in no fixed layout
            notes = ["%s: %s" % ((_cell(head[i]) or "Note"), _cell(r[i])) for i in range(15, min(len(r), 26)) if _cell(r[i])]
            v["notes"] = " · ".join(notes)
            con.execute("INSERT OR REPLACE INTO ops_vehicles(%s, created_by, created_at, updated_by, updated_at) VALUES(%s)"
                        % (", ".join(VEHICLE_FIELDS), ", ".join("?" * (len(VEHICLE_FIELDS) + 4))),
                        [v[f] for f in VEHICLE_FIELDS] + [who, now, who, now])
            nv += 1

        # columns O and P (workshop days, replacement type) are formulas, worked out on read
        cols = {0: "plate", 1: "swvl_no", 2: "vehicle_status", 3: "vehicle_type", 4: "replacement", 5: "supplier",
                6: "replacement_date", 7: "emirate", 8: "driver", 9: "contact", 10: "status", 11: "return_date",
                12: "workshop_date", 13: "perm_return_date", 16: "issue", 17: "issue_date", 18: "remarks"}
        nr = 0
        for r in list(wb["Replacement Vehicles"].iter_rows(values_only=True))[1:]:
            if not r or not _cell(r[0]):
                continue
            x = {f: _cell(r[i] if i < len(r) else None, f) for i, f in cols.items()}
            x["contact"] = _phone(x["contact"])
            if x["status"] == "Permanent Vehicle In Workhop - Replacement Vehicle Also Break Down":
                x["status"] = STATUSES[-1]
            for f in DATES & set(x):
                if x[f] and not _iso(x[f]):   # free text in a date column: keep it in the remarks
                    x["remarks"] = ("%s %s: %s" % (x["remarks"], f.replace("_", " "), x[f])).strip()
                    x[f] = ""
            con.execute("INSERT INTO ops_replacements(%s, created_by, created_at, updated_by, updated_at) VALUES(%s)"
                        % (", ".join(REPLACEMENT_FIELDS), ", ".join("?" * (len(REPLACEMENT_FIELDS) + 4))),
                        [x[f] for f in REPLACEMENT_FIELDS] + [who, now, who, now])
            nr += 1
    wb.close()
    return nv, nr


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "import":
        n_v, n_r = import_xlsx(sys.argv[2], replace="--replace" in sys.argv)
        print("Imported %d permanent vehicles and %d replacement entries." % (n_v, n_r))
    else:
        print(__doc__)
